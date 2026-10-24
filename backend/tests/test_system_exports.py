from io import BytesIO

from openpyxl import load_workbook

from app.modules.system_tools.service import ReportData, _pdf_text, build_xlsx


def test_xlsx_export_is_valid_and_blocks_formula_injection():
    report = ReportData(
        title="گزارش تست",
        columns=["محصول", "تنوع"],
        rows=[["=HYPERLINK(\"https://example.com\")", "مشکی / بزرگ"]],
    )

    payload = build_xlsx(report)
    assert payload.startswith(b"PK")

    wb = load_workbook(BytesIO(payload), data_only=False)
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True
    assert ws["A2"].value.startswith("'=")
    assert ws["B2"].value == "مشکی / بزرگ"


def test_pdf_text_escapes_reportlab_markup():
    value = _pdf_text("داریو & <تنوع>")
    assert "&amp;" in value
    assert "&lt;" in value
    assert "&gt;" in value
